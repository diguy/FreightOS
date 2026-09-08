from dataclasses import asdict, dataclass
from pathlib import Path
import re
from typing import Iterable


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md", ".markdown"}


@dataclass(frozen=True)
class DocumentMetadata:
    carrier: str
    topic: str
    document_name: str
    source_name: str
    source_type: str
    query_date: str
    effective_date: str = ""
    version: str = ""
    region: str = ""
    applicable_scope: str = ""
    knowledge_status: str = "待确认"
    file_name: str = ""

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass(frozen=True)
class KnowledgeChunk:
    chunk_id: str
    text: str
    metadata: DocumentMetadata

    def to_dict(self) -> dict:
        return {
            "chunk_id": self.chunk_id,
            "text": self.text,
            "metadata": self.metadata.to_dict(),
        }


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise ValueError(f"Unsupported document type: {path.suffix}")


def clean_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"^[ \t]*页面[ \t]*\d+[ \t]*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"^[ \t]*第?[ \t]*\d+[ \t]*页[ \t]*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = [line.strip() for line in text.splitlines()]
    deduplicated: list[str] = []
    previous = None
    for line in lines:
        if not line:
            if deduplicated and deduplicated[-1] != "":
                deduplicated.append("")
            continue
        if line == previous:
            continue
        deduplicated.append(line)
        previous = line
    return "\n".join(deduplicated).strip()


def _find_topic(path: Path) -> str:
    name = path.stem
    for keyword, topic in (
        ("禁", "禁寄和限寄"),
        ("运费", "计费规则"),
        ("报价", "计费规则"),
        ("时效", "配送时效"),
        ("服务", "服务条款"),
        ("操作", "操作指南"),
        ("FAQ", "FAQ"),
    ):
        if keyword in name:
            return topic
    return "通用物流规则"


def infer_metadata(path: Path, query_date: str) -> DocumentMetadata:
    name = path.stem
    carrier = "通用"
    for candidate in ("顺丰", "圆通", "京东", "中通", "韵达", "申通"):
        if candidate in name:
            carrier = candidate
            break
    return DocumentMetadata(
        carrier=carrier,
        topic=_find_topic(path),
        document_name=name,
        source_name=name,
        source_type="本地资料",
        query_date=query_date,
        file_name=path.name,
    )


def chunk_text(text: str, max_chars: int = 1200, overlap: int = 120) -> list[str]:
    if max_chars <= overlap:
        raise ValueError("max_chars must be greater than overlap")
    paragraphs = [item.strip() for item in re.split(r"\n{2,}", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) <= max_chars and len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            chunks.append(current)
        if len(paragraph) <= max_chars:
            current = paragraph
            continue
        start = 0
        while start < len(paragraph):
            end = min(start + max_chars, len(paragraph))
            chunks.append(paragraph[start:end].strip())
            if end == len(paragraph):
                break
            start = end - overlap
        current = ""
    if current:
        chunks.append(current)
    return chunks


def build_chunks(
    path: Path,
    query_date: str,
    *,
    max_chars: int = 1200,
    overlap: int = 120,
) -> list[KnowledgeChunk]:
    metadata = infer_metadata(path, query_date)
    text = clean_text(extract_text(path))
    return [
        KnowledgeChunk(
            chunk_id=f"{path.stem}-{index:04d}",
            text=chunk,
            metadata=metadata,
        )
        for index, chunk in enumerate(chunk_text(text, max_chars, overlap), start=1)
    ]


def build_corpus(
    input_dir: Path,
    query_date: str,
    *,
    max_chars: int = 1200,
    overlap: int = 120,
) -> list[KnowledgeChunk]:
    paths = sorted(
        path
        for path in input_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )
    chunks: list[KnowledgeChunk] = []
    for path in paths:
        chunks.extend(build_chunks(path, query_date, max_chars=max_chars, overlap=overlap))
    return chunks


def write_jsonl(chunks: Iterable[KnowledgeChunk], output_path: Path) -> int:
    import json

    output_path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with output_path.open("w", encoding="utf-8") as stream:
        for chunk in chunks:
            stream.write(json.dumps(chunk.to_dict(), ensure_ascii=False) + "\n")
            count += 1
    return count

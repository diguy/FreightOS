from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(r"F:\agent学习路径")
DOCX_PATH = OUT_DIR / "电商物流智能客服_项目交接文档_当前进度与下一步_2026-09-07.docx"
TXT_PATH = OUT_DIR / "新窗口项目交接提示词_2026-09-07.txt"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    borders_el = tbl_pr.first_child_found_in("w:tblBorders")
    if borders_el is None:
        borders_el = OxmlElement("w:tblBorders")
        tbl_pr.append(borders_el)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders_el.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders_el.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def set_cell(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    borders(table)
    for i, header in enumerate(headers):
        set_cell(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[i], "1F4E78")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], str(value))
            if row_index % 2:
                shade(cells[i], "F3F6F9")
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("电商物流智能客服项目交接文档")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("当前进度与下一步计划 2026年9月7日").bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.add_run("文档目的：").bold = True
    p.add_run("向下一窗口完整交接当前代码、Dify 工作流、验证结果和后续实施边界。本文结合项目设计文档与既有项目计划，当前实际状态以本交接文档、代码测试和最新 Dify 实测结果为准。")

    doc.add_heading("一 当前结论", level=1)
    doc.add_paragraph("意图识别与 Python 会话编排链路已经完成可运行闭环，并已通过本机 Dify 的真实端到端验证。当前最新 DSL 是 v4，已修复代码节点缩进、HTTP JSON Body 结构、结构化 result_json 接收以及测试用户身份配置。")
    add_table(doc, ["项目", "当前状态", "证据"], [
        ("Dify API", "本机 Endpoint 可用，最新 API Key 已接入", "GET /v1/parameters 返回 200"),
        ("代码节点", "无 IndentationError，可正常执行", "validator 代码编译通过，Dify 实测通过"),
        ("HTTP 节点", "单项 JSON Body，result_json 映射正常", "Dify 请求已到达 /api/v1/agent/turn"),
        ("物流查询", "已通过", "tracking_query，调用 tracking 工具成功"),
        ("人工转接", "已通过", "human_transfer，缺少 content 时阻止工具调用"),
        ("后端测试", "已通过", "poetry run pytest tests -q：180 passed"),
    ], [1.25, 2.0, 3.8])

    doc.add_heading("二 已完成工作", level=1)
    doc.add_heading("2.1 Dify 意图识别与评测链路", level=2)
    add_bullets(doc, [
        "建立并迭代物流意图识别 Chatflow，当前交付文件为 物流意图识别-v3-多轮端到端验证-意图评测修正版-v4.yml。",
        "保留意图、置信度、实体、缺失槽位、action、should_call_tool 和 result_json 的稳定输出契约。",
        "修复人工转接场景：我需要人工处理地址问题，联系方式 138****0001 必须保留 content 缺失，并禁止调用工单工具。",
        "修复代码节点缩进错误，并通过本地编译和 Dify 沙箱执行验证。",
    ])
    doc.add_heading("2.2 Python 后端编排层", level=2)
    add_bullets(doc, [
        "Dify 负责当前轮意图识别；Python 负责 result_json 解析、会话状态合并、槽位校验、确认门控、用户权限和业务工具调用。",
        "新增并验证结构化 result_json 对象格式，同时兼容历史字符串格式。",
        "FastAPI /api/v1/agent/turn 已能接收 Dify HTTP 节点请求，并将结果交给 ChatService、SessionManager 和 BusinessToolExecutor。",
        "物流查询测试使用 X-User-Id: demo-user-001，已修复固定测试用户导致的订单权限拒绝问题。",
    ])
    doc.add_heading("2.3 测试与外部验证", level=2)
    add_bullets(doc, [
        "本地测试：180 passed。重点回归覆盖 Dify 输出解析、DSL validator、Agent Turn 路由、会话状态和工具门控。",
        "真实 Dify 测试：查一下 ORD1001 到哪里了 -> tracking_query，成功调用 tracking；我需要人工处理地址问题，联系方式 138****0001 -> human_transfer，missing_slots=[content]，should_call_tool=false。",
        "已将 v2、v3、v4 DSL 文件分别保留，未覆盖历史版本。最新文件已复制到 F:\\agent学习路径。",
    ])

    doc.add_heading("三 当前文件与运行约定", level=1)
    add_table(doc, ["类别", "位置或约定", "说明"], [
        ("代码仓库", "E:\\PythonProject5\\backend-lab", "FastAPI、业务服务、测试和 DSL 交付文件"),
        ("最新 DSL", "F:\\agent学习路径\\物流意图识别-v3-多轮端到端验证-意图评测修正版-v4.yml", "导入并发布该版本"),
        ("Dify Endpoint", "http://localhost/v1", "后端配置 DIFY_BASE_URL=http://localhost"),
        ("后端入口", "http://localhost:8000", "Dify Docker HTTP 节点使用 host.docker.internal:8000"),
        ("测试用户", "demo-user-001", "v4 HTTP 节点的 X-User-Id，匹配 ORD1001 归属"),
        ("验证命令", "poetry run pytest tests -q", "修改后先跑聚焦测试，再跑全量测试"),
    ], [1.25, 3.1, 2.7])

    doc.add_heading("四 尚未完成与风险", level=1)
    add_bullets(doc, [
        "当前 Dify 真实验证覆盖关键正向和边界样本，但正式 50 条意图回归仍需在最新 v4 发布版本上重新采集并出报告。",
        "当前会话存储默认仍是 InMemorySessionStore；Redis 会话存储和 TTL 生产化尚未完成切换。",
        "当前业务数据主链路仍以 SQLite 为准；MySQL 迁移、重复初始化和 tracking_events 唯一性需要按项目计划继续验收。",
        "工单状态查询模块是下一项业务模块，需补齐 ticket_status 的 service、repository、route、权限和重复查询测试。",
        "MCP 适配层已有基础测试，但正式接入业务模型和上线观测仍待后续阶段完成。",
    ])

    doc.add_heading("五 下一步执行顺序", level=1)
    add_numbered(doc, [
        "确认当前 Dify 应用仍使用 v4 DSL 和最新 API Key，使用两个核心样本各执行一次，并保存运行记录。",
        "运行 poetry run python scripts/run_formal_intent_regression.py，生成最新 50 条意图回归预测与报告；若 Dify 输出映射异常，先停止统计并修复链路。",
        "复核 intent accuracy、entity accuracy、missing_slots accuracy 和 tool_gate accuracy，重点检查人工转接边界与地址修改确认门控。",
        "按项目计划实现并测试工单状态查询模块，覆盖 ticket_no、user_id 权限、状态枚举和不存在工单。",
        "启动 logistics-mysql 和 logistics-redis，执行 /health/dependencies，确认 MySQL、Redis 和 SQLite 回退行为。",
        "完成 MySQL 迁移和 tracking_events 去重验证，再将 Redis 用于查询缓存和会话 TTL。",
        "完成 50 条回归和基础设施验收后，再进入 MCP 业务工具扩展与交付文档更新。",
    ])

    doc.add_heading("六 新窗口提示词", level=1)
    doc.add_paragraph("以下提示词可与本交接文档、项目进度计划和项目设计文档一起发送给新窗口。")
    prompt = (
        "你将接手电商物流智能客服项目。请先阅读我同时提供的三份资料：项目交接文档、项目进度与后续交付计划、项目设计文档。\n\n"
        "请以 E:\\PythonProject5\\backend-lab 为代码仓库，遵守仓库 AGENTS.md：先读相关代码和测试，保留 dirty worktree，不修改 vendor，不覆盖历史 DSL；Dify DSL 修改必须生成新的版本文件。\n\n"
        "当前状态：物流意图识别 Chatflow 的 v4 已导入并发布；本机 Dify Endpoint 是 http://localhost/v1，后端 DIFY_BASE_URL 应保持 http://localhost；FastAPI 运行在 http://localhost:8000，Dify Docker HTTP 节点调用 http://host.docker.internal:8000/api/v1/agent/turn；v4 使用 X-User-Id=demo-user-001。代码节点、HTTP JSON Body、结构化 result_json 映射和后端对象解析已修复。\n\n"
        "已验证结果：查一下 ORD1001 到哪里了 -> tracking_query，调用 tracking 工具成功；我需要人工处理地址问题，联系方式 138****0001 -> human_transfer，missing_slots=[content]，should_call_tool=false；poetry run pytest tests -q 最近为 180 passed。\n\n"
        "请从以下下一步开始：1）先检查当前 Dify v4 发布版本和后端是否可用；2）运行 scripts/run_formal_intent_regression.py，完成最新 50 条意图回归并输出准确率、实体、缺失槽位和工具门控报告；3）若回归通过，继续实现工单状态查询模块；4）按计划验收 MySQL、Redis 和 /health/dependencies；5）每一步先说明目标、修改文件和验收标准，先跑聚焦测试再跑全量测试。不要把历史文档中的旧测试数字当作当前事实，以当前代码、测试和真实 Dify 运行结果为准。"
    )
    for paragraph in prompt.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(paragraph)

    doc.add_heading("七 交接验收标准", level=1)
    add_bullets(doc, [
        "新窗口能够复现两个核心 Dify 样本，并看到后端 200 响应。",
        "正式 50 条回归报告能明确记录样本总数、各项指标、失败样本和是否允许继续扩展。",
        "工单状态查询有独立测试，越权访问不会泄露工单信息。",
        "MySQL、Redis 健康检查和数据迁移结果有命令与输出记录。",
        "每次 DSL 修改都有新版本文件、导入发布记录和对应回归结果。",
    ])

    doc.save(DOCX_PATH)
    TXT_PATH.write_text(prompt + "\n", encoding="utf-8")
    print(DOCX_PATH)
    print(TXT_PATH)


if __name__ == "__main__":
    build_docx()

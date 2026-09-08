"""Append the current acceptance handoff to the project delivery document."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


SOURCE = Path(r"F:\agent学习路径\物流知识库\电商物流智能客服_项目进度与后续交付计划.docx")
OUTPUT = Path(
    "deliverables/电商物流智能客服_项目进度与后续交付计划_验收交接版_2026-09-07.docx"
)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    document.add_paragraph()


def main() -> None:
    document = Document(SOURCE)
    document.add_page_break()
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("当前验收与新窗口交接")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("验收日期 2026年9月7日")

    document.add_paragraph(
        "本章是当前项目状态的有效交接基线。原文档中的历史测试数量、"
        "未接入服务和早期下一步计划属于过程记录；如与本章、代码和测试结果冲突，"
        "以本章和当前工作区为准。"
    )

    add_heading(document, "一 当前交付结论", 1)
    document.add_paragraph(
        "本地知识库业务闭环、真实 Dify 只读链路、RAGFlow 检索、Redis 会话存储、"
        "MySQL 业务数据持久化和生产化健康检查均已完成。项目当前可以使用真实 MySQL "
        "和 Redis 运行，订单与工单业务已从 SQLite 迁移到 MySQL。"
    )

    add_heading(document, "二 已完成模块", 1)
    add_table(
        document,
        ["模块", "当前状态", "关键证据"],
        [
            ["Dify 输出解析", "已完成", "IntentResult 契约、异常降级、工具调用门控"],
            ["会话状态", "已完成", "SessionManager、TTL、RedisSessionStore"],
            ["订单与工单", "已完成", "查询、权限、状态流转、幂等和路由测试"],
            ["知识库", "已完成", "本地知识库、RAGFlow 主检索、失败回退"],
            ["MySQL", "已完成", "MySQL 仓储、SQLite 迁移、连接池、业务 schema"],
            ["运维能力", "已完成", "依赖健康检查、MySQL 备份脚本、运维文档"],
        ],
    )

    add_heading(document, "三 完整验收结果", 1)
    add_table(
        document,
        ["检查层级", "命令或场景", "结果"],
        [
            ["单元与模块测试", "poetry run pytest tests -q", "161 passed，1 个 Starlette/httpx 弃用警告"],
            ["重点回归", "MySQL、订单、知识库、ChatService 相关测试", "32 passed"],
            ["知识库评测", "poetry run python evaluation/evaluate_knowledge.py", "5/5 passed，使用 RAGFlow"],
            ["真实 Dify 链路", "scripts/test_live_dify_chain.py", "通过，ORD1001 只读查询成功"],
            ["真实依赖检查", "GET /health/dependencies", "MySQL ok，Redis ok"],
            ["知识库 HTTP 接口", "POST /api/v1/knowledge/query", "200，返回 records 和 context"],
            ["MySQL 数据核对", "users/orders/events/tickets", "2/3/5/54"],
            ["MySQL 备份", "scripts/backup_mysql.py", "已生成有效 SQL 备份"],
        ],
    )

    add_heading(document, "四 重要修复", 1)
    add_bullet(
        document,
        "修复 MySQL tracking_events 重复写入问题：清理历史重复记录并增加唯一约束。",
    )
    add_bullet(
        document,
        "应用初始化可重复执行；重复启动后物流轨迹数量保持为 5 条。",
    )
    add_bullet(
        document,
        "备份脚本支持宿主机 mysqldump 和 Docker 容器导出，并使用 --no-tablespaces 避免额外 PROCESS 权限要求。",
    )
    add_bullet(
        document,
        "原始交付文档保留不覆盖，本文件是 2026年9月7日验收交接版。",
    )

    add_heading(document, "五 当前运行配置", 1)
    add_table(
        document,
        ["配置项", "当前值或说明"],
        [
            ["APP_DATABASE", "mysql"],
            ["SESSION_STORE", "redis"],
            ["业务 MySQL", "本地 Docker logistics-mysql，宿主机端口 3308"],
            ["业务 Redis", "本地 Docker logistics-redis，宿主机端口 6380"],
            ["RAGFlow", "已接入真实检索，作为知识库主检索后端"],
            ["敏感配置", "只保存在 .env，不写入交接文档"],
        ],
    )

    add_heading(document, "六 尚未作为完整模型指标交付的部分", 1)
    document.add_paragraph(
        "当前 evaluation/results 中的历史 Dify 输出文件只覆盖部分用例。用 "
        "formal-20-dify_outputs-2026-09-06.jsonl 运行全量评测时，基础集 35 条中有 "
        "23 条缺少预测，对抗集 15 条中有 7 条缺少预测。因此该统计只能作为历史部分样本，"
        "不能当作完整意图准确率。工具安全门控没有发现不安全调用。下一窗口应先补齐 50 条用例的"
        "最新 Dify 输出，再形成正式准确率报告。"
    )

    add_heading(document, "七 新窗口启动步骤", 1)
    for text in [
        "读取 AGENTS.md、本文件第六至八章、docs/production-operations.md、docs/session-contract.md 和当前 git status。",
        "确认 Docker 容器 logistics-mysql、logistics-redis 仍在运行。",
        "运行 poetry run pytest tests -q，确认 161 passed 或记录变化原因。",
        "运行 GET /health/dependencies，确认 MySQL 和 Redis 均为 ok。",
        "如需更新知识库，先运行知识库评测，再修改 RAGFlow 或 Dify 配置。",
        "如需修改 DSL，复制为新的版本文件，保留原 YAML，并先说明修改范围。",
        "下一项推荐工作：补齐 50 条意图评测输出，建立正式回归报告；之后再评估 MCP 适配。",
    ]:
        add_bullet(document, text)

    add_heading(document, "八 可复用验收命令", 1)
    commands = [
        "poetry run pytest tests -q",
        "poetry run python evaluation/evaluate_knowledge.py",
        "poetry run python scripts/test_live_dify_chain.py --session-id acceptance-live-20260907",
        "$env:MYSQL_BACKUP_CONTAINER='logistics-mysql'; poetry run python scripts/backup_mysql.py",
    ]
    for command in commands:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
